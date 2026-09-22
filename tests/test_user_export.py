import unittest
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from bot.user_export import USER_EXPORT_COLUMNS, build_users_workbook


class UserExportTest(unittest.TestCase):
    def test_exports_all_known_user_fields_to_formatted_workbook(self) -> None:
        user = {
            "chat_id": 42,
            "username": "ivan",
            "first_name": "=HYPERLINK(\"https://example.com\")",
            "group_name": "ИС2-261-ОБ",
            "subgroup": 1,
            "active": 1,
            "next_lesson_notifications": 0,
            "created_at": "2026-09-01T10:00:00",
            "updated_at": "2026-09-22T12:00:00",
        }

        content = build_users_workbook([user])
        workbook = load_workbook(BytesIO(content), read_only=False)
        worksheet = workbook["Пользователи"]

        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            list(USER_EXPORT_COLUMNS),
        )
        self.assertEqual(worksheet["A2"].value, 42)
        self.assertEqual(worksheet["C2"].value, user["first_name"])
        self.assertEqual(worksheet["C2"].data_type, "s")
        self.assertEqual(worksheet["D2"].value, "ИС2-261-ОБ")
        self.assertEqual(worksheet["H2"].value, datetime(2026, 9, 1, 10, 0))
        self.assertEqual(worksheet["I2"].value, datetime(2026, 9, 22, 12, 0))
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, "A1:I2")
        self.assertTrue(worksheet["A1"].font.bold)
        workbook.close()

    def test_exports_empty_user_list(self) -> None:
        content = build_users_workbook([])
        workbook = load_workbook(BytesIO(content), read_only=False)
        worksheet = workbook["Пользователи"]

        self.assertEqual(worksheet.max_row, 1)
        self.assertEqual([cell.value for cell in worksheet[1]], list(USER_EXPORT_COLUMNS))
        workbook.close()


if __name__ == "__main__":
    unittest.main()
