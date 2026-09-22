from collections.abc import Mapping, Sequence
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


USER_EXPORT_COLUMNS = (
    "chat_id",
    "username",
    "first_name",
    "group_name",
    "subgroup",
    "active",
    "next_lesson_notifications",
    "created_at",
    "updated_at",
)

USER_EXPORT_WIDTHS = (16, 24, 24, 18, 12, 12, 28, 22, 22)
DATE_COLUMNS = {"created_at", "updated_at"}


def _excel_value(column: str, value: object) -> object:
    if column in DATE_COLUMNS and isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return value
    return value


def build_users_workbook(users: Sequence[Mapping[str, object]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Пользователи"
    worksheet.append(USER_EXPORT_COLUMNS)

    for user in users:
        worksheet.append([
            _excel_value(column, user[column]) for column in USER_EXPORT_COLUMNS
        ])

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            # Telegram names are user-controlled. Keep leading "=" as text so a
            # crafted name cannot become an Excel formula when the file is opened.
            if isinstance(cell.value, str):
                cell.data_type = "s"
        row[7].number_format = "yyyy-mm-dd hh:mm:ss"
        row[8].number_format = "yyyy-mm-dd hh:mm:ss"

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24

    for index, width in enumerate(USER_EXPORT_WIDTHS, start=1):
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
